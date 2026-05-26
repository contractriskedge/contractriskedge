"""DOCX export with native Word tracked changes using python-docx.

Exports redline suggestions as DOCX files with native Word tracked
changes (<w:ins> and <w:del> elements), allowing users to review
AI-generated changes directly in Microsoft Word.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from ..redline.models import RedlineSuggestion
from .tracked_changes import TrackedChangeManager, TrackedChange

logger = logging.getLogger(__name__)


@dataclass
class DocxExportResult:
    """Result of a DOCX export operation."""

    success: bool
    file_path: Optional[str] = None
    file_bytes: Optional[bytes] = None
    suggestion_count: int = 0
    error: Optional[str] = None
    export_time_ms: float = 0.0


class DocxExporter:
    """Exports redline suggestions as DOCX with native tracked changes.

    Creates a Word document where each redline suggestion is rendered
    using native tracked changes: original text as deleted (red,
    strikethrough) and proposed text as inserted (green, underline).

    Usage:
        exporter = DocxExporter()
        result = exporter.export(suggestions, "output.docx")
        bytes_result = exporter.export_to_bytes(suggestions)
    """

    def __init__(self) -> None:
        """Initialize the DOCX exporter."""
        self._tracked_changes = TrackedChangeManager()

    def export(
        self,
        suggestions: List[RedlineSuggestion],
        output_path: str,
        title: Optional[str] = None,
        include_metadata: bool = True,
    ) -> DocxExportResult:
        """Export redline suggestions to a DOCX file.

        Args:
            suggestions: List of redline suggestions to export.
            output_path: Path to write the DOCX file.
            title: Optional document title.
            include_metadata: Whether to include suggestion metadata.

        Returns:
            DocxExportResult with the result of the export.
        """
        start_time = datetime.utcnow()

        try:
            doc = self._build_document(suggestions, title, include_metadata)
            doc.save(output_path)

            elapsed = (datetime.utcnow() - start_time).total_seconds() * 1000

            logger.info(
                "Exported %d suggestions to %s in %.0fms",
                len(suggestions), output_path, elapsed,
            )

            return DocxExportResult(
                success=True,
                file_path=output_path,
                suggestion_count=len(suggestions),
                export_time_ms=elapsed,
            )

        except Exception as exc:
            logger.error("DOCX export failed: %s", exc, exc_info=True)
            return DocxExportResult(
                success=False,
                error=str(exc),
                suggestion_count=len(suggestions),
            )

    def export_to_bytes(
        self,
        suggestions: List[RedlineSuggestion],
        title: Optional[str] = None,
        include_metadata: bool = True,
    ) -> DocxExportResult:
        """Export redline suggestions to a bytes buffer.

        Args:
            suggestions: List of redline suggestions to export.
            title: Optional document title.
            include_metadata: Whether to include suggestion metadata.

        Returns:
            DocxExportResult with file_bytes populated.
        """
        start_time = datetime.utcnow()

        try:
            doc = self._build_document(suggestions, title, include_metadata)
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            elapsed = (datetime.utcnow() - start_time).total_seconds() * 1000

            return DocxExportResult(
                success=True,
                file_bytes=buffer.getvalue(),
                suggestion_count=len(suggestions),
                export_time_ms=elapsed,
            )

        except Exception as exc:
            logger.error("DOCX bytes export failed: %s", exc, exc_info=True)
            return DocxExportResult(
                success=False,
                error=str(exc),
                suggestion_count=len(suggestions),
            )

    def _build_document(
        self,
        suggestions: List[RedlineSuggestion],
        title: Optional[str] = None,
        include_metadata: bool = True,
    ) -> Document:
        """Build a python-docx Document with tracked changes.

        Args:
            suggestions: List of redline suggestions.
            title: Optional document title.
            include_metadata: Whether to include metadata.

        Returns:
            A configured Document instance.
        """
        doc = Document()

        # Set default font
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        # Add title
        doc_title = title or "Contract Redline Suggestions"
        heading = doc.add_heading(doc_title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add metadata header
        if include_metadata:
            meta_para = doc.add_paragraph()
            meta_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = meta_para.add_run(
                f"Generated by: AI Contract Risk Analyzer\n"
                f"Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}\n"
                f"Total Suggestions: {len(suggestions)}"
            )
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            doc.add_paragraph("─" * 80)

        # Add each suggestion
        for idx, suggestion in enumerate(suggestions, 1):
            self._add_suggestion_to_doc(doc, suggestion, idx, include_metadata)

            # Add page break between suggestions (except last)
            if idx < len(suggestions):
                doc.add_page_break()

        return doc

    def _add_suggestion_to_doc(
        self,
        doc: Document,
        suggestion: RedlineSuggestion,
        index: int,
        include_metadata: bool,
    ) -> None:
        """Add a single redline suggestion to the document.

        Args:
            doc: The Document to add to.
            suggestion: The suggestion to add.
            index: The suggestion index number.
            include_metadata: Whether to include metadata.
        """
        # Suggestion header
        heading = doc.add_heading(
            f"Suggestion {index}: {suggestion.clause_type.value.replace('_', ' ').title()}",
            level=2,
        )

        # Summary info
        info = doc.add_paragraph()
        info.style = doc.styles["Normal"]
        run = info.add_run(
            f"Confidence: {suggestion.confidence:.0%} | "
            f"Risk Impact: {suggestion.risk_impact.upper()} | "
            f"Change Type: {suggestion.change_type.title()} | "
            f"Status: {suggestion.status.title()}"
        )
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        if suggestion.attorney_review_required:
            review_run = info.add_run(" | ⚠ ATTORNEY REVIEW REQUIRED")
            review_run.font.size = Pt(9)
            review_run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
            review_run.bold = True

        # Original text label
        doc.add_heading("Original Text:", level=3)
        orig_para = doc.add_paragraph()

        # Add original text as deletion tracked change
        self._tracked_changes.create_run_with_tracked_change(
            orig_para,
            suggestion.original_text,
            change_type="deletion",
        )

        # Proposed text label
        doc.add_heading("Proposed Text:", level=3)
        prop_para = doc.add_paragraph()

        # Add proposed text as insertion tracked change
        self._tracked_changes.create_run_with_tracked_change(
            prop_para,
            suggestion.proposed_text,
            change_type="insertion",
        )

        # Rationale
        doc.add_heading("Rationale:", level=3)
        rationale_para = doc.add_paragraph(suggestion.rationale)
        rationale_para.style = doc.styles["Normal"]

        # Word diff summary
        if suggestion.word_diff:
            diff_para = doc.add_paragraph()
            diff_run = diff_para.add_run(f"Diff: {suggestion.word_diff}")
            diff_run.font.size = Pt(9)
            diff_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            diff_run.italic = True

        # Metadata section
        if include_metadata and suggestion.metadata:
            doc.add_heading("Details:", level=3)
            meta = suggestion.metadata

            issues = meta.get("issues_found", [])
            if issues:
                for issue in issues:
                    issue_para = doc.add_paragraph()
                    severity = issue.get("severity", "info")
                    issue_text = issue.get("issue", "Unknown issue")
                    issue_run = issue_para.add_run(f"• [{severity.upper()}] {issue_text}")
                    issue_run.font.size = Pt(10)

            strategy = meta.get("negotiation_strategy")
            if strategy:
                doc.add_heading("Negotiation Strategy:", level=3)
                for category, items in strategy.items():
                    if items:
                        strat_para = doc.add_paragraph()
                        strat_run = strat_para.add_run(f"{category.replace('_', ' ').title()}: ")
                        strat_run.bold = True
                        strat_run.font.size = Pt(10)
                        strat_run = strat_para.add_run(", ".join(items))
                        strat_run.font.size = Pt(10)

        # Separator
        doc.add_paragraph("─" * 80)
