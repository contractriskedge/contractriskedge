#!/usr/bin/env python3
"""Convert production-fastapi-backend.md to Word .docx"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(9)

title = doc.add_heading('ContractRiskEdge \u2014 Production FastAPI Backend Implementation Blueprint', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('V1 Modular Monolith \u2014 AI-Native Contract Intelligence Platform')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs:
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(100, 116, 139)
doc.add_paragraph()

with open('docs/production-fastapi-backend.md', 'r') as f:
    content = f.read()

lines = content.split('\n')
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith('# ContractRiskEdge') or line.strip() == '---':
        i += 1; continue
    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('#### '):
        doc.add_heading(line[5:].strip(), level=4)
    elif line.strip().startswith('|') and line.strip().endswith('|'):
        tbl_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            tbl_lines.append(lines[i].strip()); i += 1
        i -= 1
        if len(tbl_lines) >= 2:
            headers = [h.strip() for h in tbl_lines[0].split('|')[1:-1]]
            rows_data = []
            for tl in tbl_lines[2:]:
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                if cells: rows_data.append(cells)
            if headers and rows_data:
                table = doc.add_table(rows=1+len(rows_data), cols=len(headers))
                table.style = 'Light Grid Accent 1'
                for j, h in enumerate(headers):
                    c = table.rows[0].cells[j]; c.text = h
                    for p in c.paragraphs:
                        for r in p.runs: r.bold = True; r.font.size = Pt(7)
                for ri, row in enumerate(rows_data):
                    for ci, val in enumerate(row):
                        if ci < len(headers):
                            c = table.rows[ri+1].cells[ci]; c.text = val
                            for p in c.paragraphs:
                                for r in p.runs: r.font.size = Pt(7)
                doc.add_paragraph()
    elif line.strip().startswith('- ') or line.strip().startswith('+ '):
        text = line.strip()[2:].strip()
        if text: doc.add_paragraph(text, style='List Bullet')
    elif line.strip().startswith('```'):
        while i < len(lines) and not lines[i].strip().startswith('```'): i += 1
    elif line.strip() and not line.strip().startswith('```'):
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line.strip())
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2]); r.bold = True
            else: p.add_run(part)
    i += 1

outpath = 'docs/production-fastapi-backend.docx'
doc.save(outpath)
print(f'Saved {outpath}')
