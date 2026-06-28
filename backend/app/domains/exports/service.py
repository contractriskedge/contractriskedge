"""Export service — generates PDF/DOCX reports for reviews, findings, and audits."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import uuid
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Optional

from sqlalchemy import select, text as sa_text
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.domains.exports.models import AuditExportArtifact, AuditExportJob
from app.domains.exports.schemas import (
    AuditExportArtifactResponse,
    AuditExportJobResponse,
    AuditExportRequest,
    ExportFormat,
    ExportRequest,
    ExportResponse,
)
from app.integrations.storage.s3 import storage_service

logger = logging.getLogger(__name__)

# ── Template Directory ─────────────────────────────────────────────

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)


@dataclass
class ExportService:
    """Generates PDF and DOCX exports for reviews, findings, and audit data."""

    session: AsyncSession
    tenant_id: str

    async def generate_review_report(
        self, review_id: str, fmt: ExportFormat = ExportFormat.PDF,
        include_findings: bool = True, include_redlines: bool = True,
        include_evidence: bool = False, include_activity: bool = False,
    ) -> tuple[bytes, str, str]:
        """Generate a review report in PDF or DOCX format.

        Returns:
            Tuple of (file_bytes, filename, content_type).
        """
        # Fetch review data
        review = await self._fetch_review(review_id)
        findings = await self._fetch_findings(review_id) if include_findings else []
        redlines = await self._fetch_redlines(review_id) if include_redlines else []
        activity = await self._fetch_activity(review_id) if include_activity else []

        if fmt == ExportFormat.PDF:
            return await self._generate_pdf(review, findings, redlines, activity)
        else:
            return await self._generate_docx(review, findings, redlines, activity)

    async def generate_audit_report(
        self, period_days: int = 30,
    ) -> tuple[bytes, str, str]:
        """Generate an audit report in PDF format."""
        # Fetch audit data
        audit_events = await self._fetch_audit_events(period_days)
        summary = await self._fetch_audit_summary(period_days)

        html = self._render_audit_html(summary, audit_events, period_days)
        pdf_bytes = self._html_to_pdf(html)
        filename = f"audit-report-{datetime.utcnow().strftime('%Y%m%d')}.pdf"
        return pdf_bytes, filename, "application/pdf"

    async def _fetch_review(self, review_id: str) -> dict:
        sql = sa_text("""
            SELECT r.*, u.filename, u.file_size, u.content_type,
                   u.created_at as upload_date
            FROM contract_reviews r
            LEFT JOIN upload_sessions u ON u.upload_id = r.upload_id
            WHERE r.review_id = :review_id AND r.tenant_id = :tenant_id
        """)
        result = await self.session.execute(sql, {
            "review_id": review_id, "tenant_id": self.tenant_id,
        })
        row = result.fetchone()
        if not row:
            raise ValueError(f"Review {review_id} not found")
        return dict(row._mapping)

    async def _fetch_findings(self, review_id: str) -> list[dict]:
        sql = sa_text("""
            SELECT * FROM review_findings
            WHERE review_id = :review_id AND tenant_id = :tenant_id
            ORDER BY
                CASE severity
                    WHEN 'critical' THEN 1 WHEN 'high' THEN 2
                    WHEN 'medium' THEN 3 WHEN 'low' THEN 4 ELSE 5
                END
        """)
        result = await self.session.execute(sql, {
            "review_id": review_id, "tenant_id": self.tenant_id,
        })
        return [dict(r._mapping) for r in result.fetchall()]

    async def _fetch_redlines(self, review_id: str) -> list[dict]:
        sql = sa_text("""
            SELECT * FROM review_redlines
            WHERE review_id = :review_id AND tenant_id = :tenant_id
            ORDER BY created_at
        """)
        result = await self.session.execute(sql, {
            "review_id": review_id, "tenant_id": self.tenant_id,
        })
        return [dict(r._mapping) for r in result.fetchall()]

    async def _fetch_activity(self, review_id: str) -> list[dict]:
        sql = sa_text("""
            SELECT * FROM review_status_history
            WHERE review_id = :review_id AND tenant_id = :tenant_id
            ORDER BY created_at DESC
            LIMIT 50
        """)
        result = await self.session.execute(sql, {
            "review_id": review_id, "tenant_id": self.tenant_id,
        })
        return [dict(r._mapping) for r in result.fetchall()]

    async def _fetch_audit_events(self, period_days: int) -> list[dict]:
        sql = sa_text("""
            SELECT event_id, event_type, entity_type, entity_id,
                   actor_id, change_summary, correlation_id, source,
                   metadata, created_at
            FROM governance_audit_events
            WHERE tenant_id = :tenant_id
              AND created_at > NOW() - make_interval(days => :period_days)
            ORDER BY created_at DESC
            LIMIT 500
        """)
        result = await self.session.execute(sql, {
            "tenant_id": self.tenant_id,
            "period_days": period_days,
        })
        return [dict(r._mapping) for r in result.fetchall()]

    async def _fetch_audit_summary(self, period_days: int) -> dict:
        sql = sa_text("""
            SELECT
                COUNT(*)::int AS total_events,
                COUNT(DISTINCT actor_id)::int AS unique_actors,
                COUNT(DISTINCT event_type)::int AS unique_types
            FROM governance_audit_events
            WHERE tenant_id = :tenant_id
              AND created_at > NOW() - make_interval(days => :period_days)
        """)
        result = await self.session.execute(sql, {
            "tenant_id": self.tenant_id,
            "period_days": period_days,
        })
        return dict(result.fetchone()._mapping)

    async def generate_immutable_audit_export(
        self,
        request: AuditExportRequest,
        actor_id: str,
        actor_role: Optional[str] = None,
    ) -> AuditExportJobResponse:
        """Create an immutable, append-only audit export job and upload artifact."""
        filter_params = {
            k: v for k, v in request.model_dump().items()
            if k not in ("output_format", "export_reason", "request_id") and v is not None
        }
        job = AuditExportJob(
            tenant_id=uuid.UUID(self.tenant_id),
            actor_id=actor_id,
            actor_role=actor_role,
            status="pending",
            output_format=request.output_format.value,
            filter_params=filter_params,
            export_reason=request.export_reason,
            request_id=request.request_id,
            chain_of_custody={
                "requested_by": actor_id,
                "requested_role": actor_role,
                "requested_at": datetime.utcnow().isoformat() + "Z",
                "export_reason": request.export_reason,
                "request_id": request.request_id,
                "filters": filter_params,
                "output_format": request.output_format.value,
                "source": "audit_export_service",
            },
        )
        self.session.add(job)
        await self.session.flush()

        try:
            events = await self._fetch_audit_events_for_export(request)
            if request.output_format == ExportFormat.JSONL:
                file_bytes = self.serialize_audit_events_jsonl(events)
                content_type = "application/x-ndjson"
            else:
                file_bytes = self.serialize_audit_events_csv(events)
                content_type = "text/csv"

            checksum = self._compute_sha256(file_bytes)
            filename = f"audit-export-{job.job_id}.{request.output_format.value}"
            storage_key = storage_service.build_object_key(self.tenant_id, filename)
            await storage_service.upload_fileobj(
                settings.s3_bucket,
                storage_key,
                file_bytes,
                content_type,
                metadata={
                    "job_id": str(job.job_id),
                    "tenant_id": self.tenant_id,
                    "export_reason": request.export_reason or "",
                },
            )

            artifact = AuditExportArtifact(
                job_id=job.job_id,
                tenant_id=uuid.UUID(self.tenant_id),
                filename=filename,
                content_type=content_type,
                content_length=len(file_bytes),
                sha256_hash=checksum,
                storage_key=storage_key,
                artifact_metadata={
                    "export_reason": request.export_reason,
                    "request_id": request.request_id,
                },
            )
            self.session.add(artifact)
            await self.session.flush()

            manifest = {
                "job_id": str(job.job_id),
                "tenant_id": self.tenant_id,
                "actor_id": actor_id,
                "actor_role": actor_role,
                "output_format": request.output_format.value,
                "filters": filter_params,
                "artifact_count": 1,
                "artifact_hashes": [checksum],
                "created_at": job.chain_of_custody["requested_at"],
                "export_reason": request.export_reason,
            }
            manifest_json = json.dumps(manifest, sort_keys=True)
            manifest_hash = self._compute_sha256(manifest_json.encode("utf-8"))
            manifest_signature = self._sign_manifest(manifest_json)

            job.artifact_count = 1
            job.checksum = checksum
            job.manifest_hash = manifest_hash
            job.manifest_signature = manifest_signature
            job.chain_of_custody = {
                **job.chain_of_custody,
                "manifest_hash": manifest_hash,
                "manifest_signature": manifest_signature,
                "artifact_hashes": [checksum],
            }
            job.status = "completed"
            job.completed_at = datetime.utcnow()
            await self.session.flush()

            from app.domains.audit.recorder import AuditRecorder
            recorder = AuditRecorder(self.session, self.tenant_id)
            await recorder.record(
                event_type="export.generated",
                actor_id=actor_id,
                actor_role=actor_role,
                description=f"Immutable audit export ({request.output_format.value}) — {len(events)} events",
                status="success",
                severity="info",
                entity_type="audit_export",
                entity_id=str(job.job_id),
                source="audit_export_service",
                metadata={
                    "job_id": str(job.job_id),
                    "output_format": request.output_format.value,
                    "event_count": len(events),
                    "filters": filter_params,
                },
            )
        except Exception as exc:
            job.status = "failed"
            job.failure_reason = str(exc)
            await self.session.flush()
            logger.exception("Audit export job failed: %s", exc)

        return await self.get_audit_export_job(str(job.job_id))

    async def _fetch_audit_events_for_export(self, request: AuditExportRequest) -> list[dict]:
        conditions = ["tenant_id = :tenant_id"]
        bind = {"tenant_id": self.tenant_id}
        if request.event_type:
            conditions.append("event_type = :event_type")
            bind["event_type"] = request.event_type
        if request.entity_type:
            conditions.append("entity_type = :entity_type")
            bind["entity_type"] = request.entity_type
        if request.entity_id:
            conditions.append("entity_id = :entity_id")
            bind["entity_id"] = request.entity_id
        if request.actor_id:
            conditions.append("actor_id = :actor_id")
            bind["actor_id"] = request.actor_id
        if request.from_date:
            conditions.append("created_at >= :from_date")
            bind["from_date"] = request.from_date
        if request.to_date:
            conditions.append("created_at <= :to_date")
            bind["to_date"] = request.to_date

        sql = sa_text(f"""
            SELECT event_id, event_type, entity_type, entity_id,
                   actor_id, actor_role, previous_state, new_state,
                   change_summary, correlation_id, request_id, source,
                   metadata, created_at
            FROM governance_audit_events
            WHERE {' AND '.join(conditions)}
            ORDER BY created_at DESC
            LIMIT 50000
        """)
        result = await self.session.execute(sql, bind)
        return [dict(r._mapping) for r in result.fetchall()]

    def serialize_audit_events_csv(self, events: list[dict]) -> bytes:
        headers = [
            "event_id", "event_type", "entity_type", "entity_id",
            "actor_id", "actor_role", "change_summary", "correlation_id",
            "request_id", "source", "status", "created_at", "metadata",
            "previous_state", "new_state",
        ]
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for event in events:
            row = {**event}
            meta = event.get("metadata") or {}
            if isinstance(meta, str):
                try:
                    meta = json.loads(meta)
                except json.JSONDecodeError:
                    meta = {}
            row["status"] = meta.get("status", "success")
            row["created_at"] = event.get("created_at").isoformat() if event.get("created_at") else ""
            row["metadata"] = json.dumps(meta, sort_keys=True)
            row["previous_state"] = json.dumps(event.get("previous_state", {}), sort_keys=True) if event.get("previous_state") is not None else ""
            row["new_state"] = json.dumps(event.get("new_state", {}), sort_keys=True) if event.get("new_state") is not None else ""
            writer.writerow(row)
        return buf.getvalue().encode("utf-8")

    def serialize_audit_events_jsonl(self, events: list[dict]) -> bytes:
        lines = []
        for event in events:
            if isinstance(event.get("created_at"), datetime):
                event["created_at"] = event["created_at"].isoformat()
            lines.append(json.dumps(event, sort_keys=True))
        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def _compute_sha256(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    def _sign_manifest(self, manifest_json: str) -> str:
        signing_key = settings.secret_key or "default-signature-key"
        signature_input = f"{signing_key}:{manifest_json}".encode("utf-8")
        return hashlib.sha256(signature_input).hexdigest()

    async def get_audit_export_job(self, job_id: str) -> AuditExportJobResponse:
        sql = sa_text("""
            SELECT * FROM audit_export_jobs
            WHERE job_id = :job_id AND tenant_id = :tenant_id
        """)
        result = await self.session.execute(sql, {
            "job_id": job_id,
            "tenant_id": self.tenant_id,
        })
        row = result.fetchone()
        if not row:
            raise ValueError(f"Audit export job {job_id} not found")
        job = dict(row._mapping)
        artifacts = await self._fetch_audit_export_artifacts(job_id)
        return AuditExportJobResponse(
            job_id=str(job["job_id"]),
            status=job["status"],
            output_format=ExportFormat(job["output_format"]),
            filter_params=job["filter_params"],
            export_reason=job["export_reason"],
            request_id=job["request_id"],
            artifact_count=job["artifact_count"],
            checksum=job["checksum"],
            manifest_hash=job["manifest_hash"],
            manifest_signature=job["manifest_signature"],
            chain_of_custody=job["chain_of_custody"],
            artifacts=artifacts,
            created_at=job["created_at"],
            completed_at=job["completed_at"],
            failure_reason=job["failure_reason"],
        )

    async def _fetch_audit_export_artifacts(self, job_id: str) -> list[AuditExportArtifactResponse]:
        sql = sa_text("""
            SELECT * FROM audit_export_artifacts
            WHERE job_id = :job_id AND tenant_id = :tenant_id
        """)
        result = await self.session.execute(sql, {
            "job_id": job_id,
            "tenant_id": self.tenant_id,
        })
        items = [dict(r._mapping) for r in result.fetchall()]
        return [
            AuditExportArtifactResponse(
                artifact_id=str(item["artifact_id"]),
                filename=item["filename"],
                content_type=item["content_type"],
                content_length=item["content_length"],
                sha256_hash=item["sha256_hash"],
                storage_key=item["storage_key"],
                metadata=item["metadata"],
                created_at=item["created_at"],
            )
            for item in items
        ]

    async def list_audit_export_jobs(self, page: int = 1, page_size: int = 50) -> list[AuditExportJobResponse]:
        offset = (page - 1) * page_size
        sql = sa_text("""
            SELECT * FROM audit_export_jobs
            WHERE tenant_id = :tenant_id
            ORDER BY created_at DESC
            LIMIT :limit OFFSET :offset
        """)
        result = await self.session.execute(sql, {
            "tenant_id": self.tenant_id,
            "limit": page_size,
            "offset": offset,
        })
        jobs = [dict(r._mapping) for r in result.fetchall()]
        return [await self.get_audit_export_job(str(job["job_id"])) for job in jobs]

    async def download_audit_export_artifact(self, job_id: str, artifact_id: str) -> tuple[bytes, str, str]:
        sql = sa_text("""
            SELECT filename, content_type, storage_key
            FROM audit_export_artifacts
            WHERE artifact_id = :artifact_id
              AND job_id = :job_id
              AND tenant_id = :tenant_id
        """)
        result = await self.session.execute(sql, {
            "artifact_id": artifact_id,
            "job_id": job_id,
            "tenant_id": self.tenant_id,
        })
        row = result.fetchone()
        if not row:
            raise ValueError("Audit export artifact not found")
        artifact = dict(row._mapping)
        if not artifact["storage_key"]:
            raise ValueError("Audit export artifact storage key missing")
        file_bytes = await storage_service.download_fileobj(settings.s3_bucket, artifact["storage_key"])
        return file_bytes, artifact["content_type"], artifact["filename"]

    async def _generate_pdf(
        self, review: dict, findings: list[dict],
        redlines: list[dict], activity: list[dict],
    ) -> tuple[bytes, str, str]:
        """Generate a PDF review report using WeasyPrint."""
        html = self._render_review_html(review, findings, redlines, activity)
        pdf_bytes = self._html_to_pdf(html)
        filename = f"review-report-{str(review.get('review_id', 'unknown'))[:8]}.pdf"
        return pdf_bytes, filename, "application/pdf"

    async def _generate_docx(
        self, review: dict, findings: list[dict],
        redlines: list[dict], activity: list[dict],
    ) -> tuple[bytes, str, str]:
        """Generate a DOCX review report."""
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            logger.error("python-docx not installed. Install with: pip install python-docx")
            raise

        doc = Document()

        # Title
        title = doc.add_heading("Contract Review Report", level=0)
        doc.add_paragraph(f"Review ID: {str(review.get('review_id', 'unknown'))[:12]}...")
        doc.add_paragraph(f"Status: {review.get('status', 'unknown')}")
        doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
        doc.add_paragraph("")

        # Findings section
        if findings:
            doc.add_heading("Findings", level=1)
            for f in findings:
                doc.add_heading(f.get("title", "Untitled"), level=2)
                doc.add_paragraph(f"Severity: {f.get('severity', 'unknown')}")
                p = doc.add_paragraph(f.get("description", ""))
                if f.get("recommendation"):
                    doc.add_paragraph(f"Recommendation: {f['recommendation']}")

        # Redlines section
        if redlines:
            doc.add_heading("Redline Suggestions", level=1)
            for r in redlines:
                doc.add_heading(f"Clause: {r.get('clause_type', 'unknown')}", level=2)
                doc.add_paragraph(f"Status: {r.get('status', 'proposed')}")
                doc.add_paragraph("Original:")
                doc.add_paragraph(r.get("original_text", ""))
                doc.add_paragraph("Proposed:")
                doc.add_paragraph(r.get("proposed_text", ""))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        filename = f"review-report-{str(review.get('review_id', 'unknown'))[:8]}.docx"
        return buf.getvalue(), filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def _render_review_html(
        self, review: dict, findings: list[dict],
        redlines: list[dict], activity: list[dict],
    ) -> str:
        """Render review report as HTML for PDF conversion."""
        severity_color = {
            "critical": "#dc2626", "high": "#ea580c",
            "medium": "#d97706", "low": "#6b7280", "info": "#3b82f6",
        }

        findings_html = ""
        for f in findings:
            color = severity_color.get(f.get("severity", ""), "#6b7280")
            findings_html += f"""
            <div class="finding" style="border-left: 4px solid {color};">
                <div class="finding-header">
                    <span class="severity" style="background: {color}; color: white; padding: 2px 8px; border-radius: 4px; font-size: 11px;">
                        {f.get('severity', 'unknown').upper()}
                    </span>
                    <h3>{f.get('title', 'Untitled')}</h3>
                </div>
                <p>{f.get('description', '')}</p>
                {f'<p class="recommendation"><strong>Recommendation:</strong> {f["recommendation"]}</p>' if f.get('recommendation') else ''}
            </div>
            """

        redlines_html = ""
        for r in redlines:
            redlines_html += f"""
            <div class="redline">
                <h3>Clause: {r.get('clause_type', 'unknown')}</h3>
                <div class="diff">
                    <div class="original"><strong>Original:</strong><br>{r.get('original_text', '')}</div>
                    <div class="proposed"><strong>Proposed:</strong><br>{r.get('proposed_text', '')}</div>
                </div>
                {f'<p class="rationale"><em>{r["rationale"]}</em></p>' if r.get('rationale') else ''}
            </div>
            """

        activity_html = ""
        for a in activity[:20]:
            activity_html += f"""
            <tr>
                <td>{a.get('from_status', '')} → {a.get('to_status', '')}</td>
                <td>{a.get('changed_by', '')}</td>
                <td>{a.get('created_at', '').strftime('%Y-%m-%d %H:%M') if hasattr(a.get('created_at'), 'strftime') else a.get('created_at', '')}</td>
            </tr>
            """

        return f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8">
<style>
    body {{ font-family: 'Helvetica', 'Arial', sans-serif; font-size: 11pt; color: #1f2937; margin: 40px; }}
    h1 {{ color: #1e3a5f; border-bottom: 2px solid #1e3a5f; padding-bottom: 8px; }}
    h2 {{ color: #374151; margin-top: 24px; }}
    .meta {{ background: #f3f4f6; padding: 16px; border-radius: 8px; margin: 16px 0; }}
    .meta p {{ margin: 4px 0; font-size: 10pt; }}
    .finding {{ margin: 16px 0; padding: 12px 16px; background: #f9fafb; border-radius: 8px; }}
    .finding-header {{ display: flex; align-items: center; gap: 12px; }}
    .finding h3 {{ margin: 0; font-size: 12pt; }}
    .finding p {{ margin: 8px 0 0 0; font-size: 10pt; color: #4b5563; }}
    .recommendation {{ color: #2563eb; }}
    .redline {{ margin: 16px 0; padding: 12px; background: #f9fafb; border-radius: 8px; }}
    .diff {{ display: flex; gap: 16px; margin: 8px 0; }}
    .original, .proposed {{ flex: 1; padding: 8px; border-radius: 4px; font-size: 10pt; }}
    .original {{ background: #fef2f2; border: 1px solid #fecaca; }}
    .proposed {{ background: #f0fdf4; border: 1px solid #bbf7d0; }}
    .rationale {{ color: #6b7280; font-size: 10pt; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 10pt; }}
    th, td {{ padding: 8px; text-align: left; border-bottom: 1px solid #e5e7eb; }}
    th {{ background: #f3f4f6; font-weight: 600; }}
    .footer {{ margin-top: 32px; padding-top: 16px; border-top: 1px solid #e5e7eb; font-size: 9pt; color: #9ca3af; text-align: center; }}
</style></head>
<body>
    <h1>Contract Review Report</h1>
    <div class="meta">
        <p><strong>Review ID:</strong> {str(review.get('review_id', 'N/A'))[:12]}...</p>
        <p><strong>Status:</strong> {review.get('status', 'N/A')}</p>
        <p><strong>Risk Score:</strong> {review.get('risk_score', 'N/A')}</p>
        <p><strong>Findings:</strong> {len(findings)}</p>
        <p><strong>Redlines:</strong> {len(redlines)}</p>
        <p><strong>Generated:</strong> {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}</p>
    </div>

    <h2>Findings ({len(findings)})</h2>
    {findings_html or '<p>No findings identified.</p>'}

    <h2>Redline Suggestions ({len(redlines)})</h2>
    {redlines_html or '<p>No redline suggestions.</p>'}

    {f'''
    <h2>Activity Timeline</h2>
    <table><tr><th>Change</th><th>By</th><th>Date</th></tr>{activity_html}</table>
    ''' if activity else ''}

    <div class="footer">ContractRiskEdge — Enterprise Contract Review Platform</div>
</body></html>"""

    def _render_audit_html(self, summary: dict, events: list[dict], period_days: int) -> str:
        events_html = ""
        for e in events[:200]:
            events_html += f"""
            <tr>
                <td>{e.get('event_type', '')}</td>
                <td>{e.get('action', '')}</td>
                <td>{e.get('resource_type', '')}</td>
                <td>{e.get('actor_id', '')}</td>
                <td>{e.get('created_at', '').strftime('%Y-%m-%d %H:%M') if hasattr(e.get('created_at'), 'strftime') else e.get('created_at', '')}</td>
            </tr>
            """

        return f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8">
<style>
    body {{ font-family: 'Helvetica', 'Arial', sans-serif; font-size: 10pt; color: #1f2937; margin: 40px; }}
    h1 {{ color: #1e3a5f; border-bottom: 2px solid #1e3a5f; }}
    .summary {{ display: flex; gap: 16px; margin: 16px 0; }}
    .stat {{ flex: 1; background: #f3f4f6; padding: 16px; border-radius: 8px; text-align: center; }}
    .stat-value {{ font-size: 24pt; font-weight: bold; color: #1e3a5f; }}
    .stat-label {{ font-size: 9pt; color: #6b7280; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 9pt; }}
    th, td {{ padding: 6px 8px; text-align: left; border-bottom: 1px solid #e5e7eb; }}
    th {{ background: #f3f4f6; }}
    .footer {{ margin-top: 32px; padding-top: 16px; border-top: 1px solid #e5e7eb; font-size: 8pt; color: #9ca3af; text-align: center; }}
</style></head>
<body>
    <h1>Audit Report</h1>
    <p>Period: Last {period_days} days | Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}</p>
    <div class="summary">
        <div class="stat"><div class="stat-value">{summary.get('total_events', 0)}</div><div class="stat-label">Total Events</div></div>
        <div class="stat"><div class="stat-value">{summary.get('unique_actors', 0)}</div><div class="stat-label">Unique Actors</div></div>
        <div class="stat"><div class="stat-value">{summary.get('unique_types', 0)}</div><div class="stat-label">Event Types</div></div>
    </div>
    <h2>Events ({len(events)})</h2>
    <table><tr><th>Type</th><th>Action</th><th>Resource</th><th>Actor</th><th>Date</th></tr>{events_html}</table>
    <div class="footer">ContractRiskEdge — Enterprise Contract Review Platform</div>
</body></html>"""

    def _html_to_pdf(self, html: str) -> bytes:
        """Convert HTML to PDF using WeasyPrint."""
        try:
            from weasyprint import HTML as WeasyPrintHTML
            pdf_bytes = WeasyPrintHTML(string=html).write_pdf()
            return pdf_bytes
        except Exception as exc:
            logger.error("PDF generation failed: %s", exc)
            # Fallback: return HTML as bytes if PDF fails
            return html.encode("utf-8")
