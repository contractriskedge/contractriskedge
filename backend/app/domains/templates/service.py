"""Template Library service — business logic for CLM authoring."""

from __future__ import annotations

import json
import logging
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Optional

from sqlalchemy.ext.asyncio import AsyncSession

from .models import (
    ContractTemplate, TemplateVersion, TemplateVariable,
    TemplateCategory, GeneratedContract, TemplateClause, TemplateClauseRef,
    TemplatePackage, TemplatePackageItem,
)
from .repository import TemplateRepository
from .schemas import (
    TemplateCreate, TemplateUpdate, TemplateVersionCreate,
    TemplateCategoryCreate, TemplateCategoryUpdate,
    TemplateClauseCreate, TemplateClauseUpdate,
    GenerateContractRequest,
    TemplateListItem, TemplateDetail, TemplateVersionResponse,
    TemplateVariableResponse, TemplateCategoryResponse,
    TemplateClauseResponse,
    GenerateContractResponse, TemplateMetrics,
    TemplateValidationResult, ValidationIssue,
    TemplateDependencyInfo, ClauseAnalytics, ClauseDependencyInfo,
    TemplatePackageCreate, TemplatePackageUpdate, TemplatePackageResponse,
    PackageItemCreate, PackageItemResponse,
)

logger = logging.getLogger(__name__)


class TemplateService:
    """Enterprise template library service."""

    def __init__(self, repo: TemplateRepository, actor_id: str):
        self.repo = repo
        self.actor_id = actor_id
        self._session = repo.session
        self._pending_ingestion: tuple[str, str, str] | None = None

    def pop_pending_ingestion(self) -> tuple[str, str, str] | None:
        """Upload/tenant/user tuple to dispatch after HTTP transaction commits."""
        pending = self._pending_ingestion
        self._pending_ingestion = None
        return pending

    # ── Categories ────────────────────────────────────────────────

    async def list_categories(self) -> list[TemplateCategoryResponse]:
        categories = await self.repo.list_categories()
        counts = await self.repo.get_category_template_count()
        return [
            TemplateCategoryResponse(
                id=c.id,
                name=c.name,
                slug=c.slug,
                description=c.description,
                icon=c.icon,
                display_order=c.display_order,
                is_active=c.is_active,
                template_count=counts.get(c.id, 0),
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in categories
        ]

    async def create_category(
        self, body: TemplateCategoryCreate,
    ) -> TemplateCategoryResponse:
        cat = TemplateCategory(
            tenant_id=self.repo.tenant_id,
            name=body.name,
            slug=body.slug,
            description=body.description,
            icon=body.icon,
            display_order=body.display_order,
        )
        cat = await self.repo.create_category(cat)
        return TemplateCategoryResponse(
            id=cat.id, name=cat.name, slug=cat.slug,
            description=cat.description, icon=cat.icon,
            display_order=cat.display_order, is_active=cat.is_active,
        )

    async def update_category(
        self, category_id: str, body: TemplateCategoryUpdate,
    ) -> Optional[TemplateCategoryResponse]:
        cat = await self.repo.get_category(category_id)
        if not cat:
            return None
        if body.name is not None:
            cat.name = body.name
        if body.slug is not None:
            cat.slug = body.slug
        if body.description is not None:
            cat.description = body.description
        if body.icon is not None:
            cat.icon = body.icon
        if body.display_order is not None:
            cat.display_order = body.display_order
        if body.is_active is not None:
            cat.is_active = body.is_active
        cat = await self.repo.update_category(cat)
        return TemplateCategoryResponse(
            id=cat.id, name=cat.name, slug=cat.slug,
            description=cat.description, icon=cat.icon,
            display_order=cat.display_order, is_active=cat.is_active,
        )

    async def delete_category(self, category_id: str) -> bool:
        return await self.repo.delete_category(category_id)

    # ── Templates ─────────────────────────────────────────────────

    async def list_templates(
        self,
        status: Optional[str] = None,
        category_id: Optional[str] = None,
        search: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[list[TemplateListItem], int]:
        templates, total = await self.repo.list_templates(
            status=status, category_id=category_id,
            search=search, page=page, page_size=page_size,
        )
        favorite_ids = set()
        try:
            favorite_ids = await self.repo.get_favorite_ids(self.actor_id)
        except Exception:
            pass

        items = []
        for t in templates:
            current_ver = t.versions[0] if t.versions else None
            items.append(TemplateListItem(
                id=t.id,
                name=t.name,
                description=t.description,
                category_id=t.category_id,
                category_name=t.category.name if t.category else None,
                tags=list(t.tags) if isinstance(t.tags, list) else [],
                owner=t.owner,
                department=t.department,
                business_unit=t.business_unit,
                status=t.status,
                current_version_number=current_ver.version_number if current_ver else None,
                current_version_label=current_ver.label if current_ver else None,
                usage_count=t.usage_count,
                is_favorite=t.id in favorite_ids,
                default_workflow=t.default_workflow,
                created_by=t.created_by,
                created_at=t.created_at.isoformat() if t.created_at else None,
                updated_by=t.updated_by,
                updated_at=t.updated_at.isoformat() if t.updated_at else None,
            ))
        return items, total

    @staticmethod
    def _ensure_variables_list(variables: Any) -> list:
        """Ensure variables is a list of dicts, handling JSON string storage."""
        if isinstance(variables, str):
            try:
                return json.loads(variables)
            except (json.JSONDecodeError, TypeError):
                return []
        if isinstance(variables, list):
            return variables
        if isinstance(variables, dict):
            return [variables]
        return []

    async def get_template(self, template_id: str) -> Optional[TemplateDetail]:
        t = await self.repo.get_template(template_id)
        if not t:
            return None

        favorite_ids = set()
        try:
            favorite_ids = await self.repo.get_favorite_ids(self.actor_id)
        except Exception:
            pass

        category_resp = None
        if t.category:
            cat = t.category
            category_resp = TemplateCategoryResponse(
                id=cat.id, name=cat.name, slug=cat.slug,
                description=cat.description, icon=cat.icon,
                display_order=cat.display_order, is_active=cat.is_active,
            )

        versions = []
        current_version = None
        for v in sorted(t.versions or [], key=lambda x: x.version_number, reverse=True):
            var_list = self._ensure_variables_list(v.variables)
            var_resp = []
            for var in var_list:
                if isinstance(var, dict):
                    try:
                        var_resp.append(TemplateVariableResponse(
                            id=var.get("id", ""),
                            key=var.get("key", "unknown"),
                            label=var.get("label", "Unknown"),
                            field_type=var.get("field_type", "text"),
                            description=var.get("description"),
                            default_value=var.get("default_value"),
                            options=var.get("options", []),
                            validation_rules=var.get("validation_rules", {}),
                            help_text=var.get("help_text"),
                            display_order=var.get("display_order", 0),
                            is_required=var.get("is_required", False),
                            variable_source=var.get("variable_source", "user_input"),
                            calculation_rule=var.get("calculation_rule"),
                        ))
                    except Exception:
                        var_resp.append(TemplateVariableResponse(
                            id="", key=var.get("key", "unknown"),
                            label=var.get("label", "Unknown"),
                        ))

            ver_resp = TemplateVersionResponse(
                id=v.id,
                template_id=v.template_id,
                version_number=v.version_number,
                label=v.label,
                change_summary=v.change_summary,
                status=v.status,
                variables=var_resp,
                clause_refs=list(v.clause_refs) if isinstance(v.clause_refs, list) else [],
                placeholder_content=v.placeholder_content,
                file_name=v.file_name,
                file_size_bytes=v.file_size_bytes,
                mime_type=v.mime_type,
                created_by=v.created_by,
                created_at=v.created_at.isoformat() if v.created_at else None,
            )
            versions.append(ver_resp)
            if v.id == t.current_version_id or (not current_version and len(versions) == 1):
                current_version = ver_resp

        return TemplateDetail(
            id=t.id,
            name=t.name,
            description=t.description,
            category=category_resp,
            tags=list(t.tags) if isinstance(t.tags, list) else [],
            owner=t.owner,
            department=t.department,
            business_unit=t.business_unit,
            status=t.status,
            current_version=current_version,
            versions=versions,
            usage_count=t.usage_count,
            is_favorite=t.id in favorite_ids,
            default_workflow=t.default_workflow,
            created_by=t.created_by,
            created_at=t.created_at.isoformat() if t.created_at else None,
            updated_by=t.updated_by,
            updated_at=t.updated_at.isoformat() if t.updated_at else None,
        )

    async def create_template(
        self, body: TemplateCreate,
    ) -> TemplateDetail:
        now = datetime.now(timezone.utc)
        template = ContractTemplate(
            tenant_id=self.repo.tenant_id,
            category_id=body.category_id,
            name=body.name,
            description=body.description,
            tags=body.tags or [],
            owner=body.owner,
            department=body.department,
            business_unit=body.business_unit,
            default_workflow=body.default_workflow,
            status="draft",
            created_by=self.actor_id,
            created_at=now,
            updated_at=now,
        )
        template = await self.repo.create_template(template)

        # Create initial version if provided
        if body.version:
            ver = await self._create_version(template.id, body.version)
            await self.repo.update_template_current_version(template.id, ver.id)

        # Log audit
        await self.repo.log_usage(
            template.id, "created", self.actor_id,
            details={"name": body.name},
        )

        return await self.get_template(template.id)

    async def update_template(
        self, template_id: str, body: TemplateUpdate,
    ) -> Optional[TemplateDetail]:
        t = await self.repo.get_template(template_id)
        if not t:
            return None
        if body.name is not None:
            t.name = body.name
        if body.description is not None:
            t.description = body.description
        if body.category_id is not None:
            t.category_id = body.category_id
        if body.tags is not None:
            t.tags = body.tags
        if body.owner is not None:
            t.owner = body.owner
        if body.department is not None:
            t.department = body.department
        if body.business_unit is not None:
            t.business_unit = body.business_unit
        if body.default_workflow is not None:
            t.default_workflow = body.default_workflow
        if body.status is not None:
            t.status = body.status
        t.updated_by = self.actor_id
        await self.repo.update_template(t)

        await self.repo.log_usage(
            template_id, "updated", self.actor_id,
            details={"changes": body.model_dump(exclude_none=True)},
        )
        return await self.get_template(template_id)

    async def delete_template(self, template_id: str) -> bool:
        result = await self.repo.delete_template(template_id)
        if result:
            await self.repo.log_usage(
                template_id, "deleted", self.actor_id,
            )
        return result

    async def archive_template(self, template_id: str) -> Optional[TemplateDetail]:
        return await self.update_template(template_id, TemplateUpdate(status="archived"))

    async def publish_template(self, template_id: str) -> Optional[TemplateDetail]:
        result = await self.update_template(template_id, TemplateUpdate(status="approved"))
        if result:
            await self.repo.log_usage(
                template_id, "published", self.actor_id,
            )
        return result

    # ── Versions ──────────────────────────────────────────────────

    async def _create_version(
        self, template_id: str, body: TemplateVersionCreate,
    ) -> TemplateVersion:
        now = datetime.now(timezone.utc)
        ver = TemplateVersion(
            tenant_id=self.repo.tenant_id,
            template_id=template_id,
            version_number=body.version_number,
            label=body.label,
            change_summary=body.change_summary,
            status="draft",
            variables=[v.model_dump() for v in body.variables] if body.variables else [],
            placeholder_content=body.placeholder_content,
            file_name=body.file_name,
            created_by=self.actor_id,
            created_at=now,
        )
        return await self.repo.create_version(ver)

    async def create_version(
        self, template_id: str, body: TemplateVersionCreate,
    ) -> Optional[TemplateDetail]:
        t = await self.repo.get_template(template_id)
        if not t:
            return None
        ver = await self._create_version(template_id, body)
        await self.repo.update_template_current_version(template_id, ver.id)
        await self.repo.log_usage(
            template_id, "version_created", self.actor_id,
            details={"version_number": body.version_number, "label": body.label},
        )
        return await self.get_template(template_id)

    async def get_version(self, version_id: str) -> Optional[TemplateVersionResponse]:
        v = await self.repo.get_version(version_id)
        if not v:
            return None
        var_list = self._ensure_variables_list(v.variables)
        var_resp = []
        for var in var_list:
            if isinstance(var, dict):
                try:
                    var_resp.append(TemplateVariableResponse(
                        id=var.get("id", ""),
                        key=var.get("key", "unknown"),
                        label=var.get("label", "Unknown"),
                        field_type=var.get("field_type", "text"),
                        description=var.get("description"),
                        default_value=var.get("default_value"),
                        options=var.get("options", []),
                        validation_rules=var.get("validation_rules", {}),
                        help_text=var.get("help_text"),
                        display_order=var.get("display_order", 0),
                        is_required=var.get("is_required", False),
                        variable_source=var.get("variable_source", "user_input"),
                        calculation_rule=var.get("calculation_rule"),
                    ))
                except Exception:
                    var_resp.append(TemplateVariableResponse(
                        id="", key=var.get("key", "unknown"),
                        label=var.get("label", "Unknown"),
                    ))
        return TemplateVersionResponse(
            id=v.id, template_id=v.template_id,
            version_number=v.version_number, label=v.label,
            change_summary=v.change_summary, status=v.status,
            variables=var_resp,
            clause_refs=list(v.clause_refs) if isinstance(v.clause_refs, list) else [],
            placeholder_content=v.placeholder_content,
            file_name=v.file_name, file_size_bytes=v.file_size_bytes,
            mime_type=v.mime_type, created_by=v.created_by,
            created_at=v.created_at.isoformat() if v.created_at else None,
        )

    # ── Favorites ─────────────────────────────────────────────────

    async def toggle_favorite(self, template_id: str) -> dict:
        is_fav = await self.repo.toggle_favorite(template_id, self.actor_id)
        return {"template_id": template_id, "is_favorite": is_fav}

    # ── Contract Generation ───────────────────────────────────────

    async def generate_contract(
        self, body: GenerateContractRequest,
    ) -> Optional[GenerateContractResponse]:
        """Generate a contract from a template.

        Features:
        - Template version pinning — body.template_version_id freezes the version
        - Idempotency key — prevents duplicate generation on retry
        - Draft support — body.status='draft' saves without finalizing
        - Preview mode — body.preview_only renders without creating records
        """
        t = await self.repo.get_template(body.template_id)
        if not t:
            return None
        if t.status != "approved" and not body.preview_only:
            raise ValueError(f"Cannot generate from template with status '{t.status}'. Only approved templates can be used.")

        now = datetime.now(timezone.utc)
        title = body.title or f"{t.name} - {now.strftime('%Y-%m-%d')}"

        # ── Idempotency check ─────────────────────────────────────
        if body.idempotency_key and not body.preview_only:
            existing = await self.repo.get_generated_contract_by_idempotency(
                body.idempotency_key,
            )
            if existing and existing.review_id and existing.status != "draft":
                # Already finalized — return it
                logger.info("Idempotency hit for key=%s, returning existing contract %s",
                            body.idempotency_key, existing.id)
                return GenerateContractResponse(
                    generated_contract_id=existing.id,
                    review_id=existing.review_id or "",
                    template_version_id=existing.template_version_id,
                    title=existing.title,
                    status=existing.status,
                    variable_values=existing.variable_values,
                    clause_selections=list(existing.clause_selections) if isinstance(existing.clause_selections, list) else [],
                    created_at=existing.created_at.isoformat() if existing.created_at else None,
                )
            elif existing and body.status != "draft":
                # Draft exists — upgrade it to finalized (don't return early)
                logger.info("Idempotency hit for key=%s, upgrading draft %s to finalized",
                            body.idempotency_key, existing.id)
                existing.status = "finalized"
                existing.title = title
                existing.variable_values = body.variable_values
                existing.clause_selections = [cs.model_dump() for cs in body.clause_selections] if body.clause_selections else []
                existing.updated_at = now
                gc = existing
                review_id = str(uuid.uuid4())
                existing.review_id = review_id
                await self.repo.session.flush()
                # Don't return — continue to create review below
            else:
                # Draft with no review — will create below
                pass

        # ── Version pinning ───────────────────────────────────────
        # If caller specified a version, use it; otherwise use current
        pinned_version_id = body.template_version_id
        current_ver = None
        if pinned_version_id:
            current_ver = await self.repo.get_version(pinned_version_id)
            if not current_ver or current_ver.template_id != body.template_id:
                raise ValueError(f"Pinned version {pinned_version_id} not found or does not belong to this template.")
        else:
            for v in (t.versions or []):
                if v.id == t.current_version_id:
                    current_ver = v
                    break
            if not current_ver and t.versions:
                current_ver = t.versions[0]

        if not current_ver:
            raise ValueError("Template has no versions")

        # Build clause selections data
        clause_selections_data = []
        if body.clause_selections:
            clause_selections_data = [cs.model_dump() for cs in body.clause_selections]

        # Substitute variables and merge selected clauses into the contract body
        generated_content = await self._assemble_generated_content(
            current_ver.placeholder_content or "",
            body.variable_values,
            clause_selections_data,
            current_ver.id,
        )

        # Preview mode — return rendered content without creating records
        if body.preview_only:
            return GenerateContractResponse(
                title=title,
                status="preview",
                variable_values=body.variable_values,
                preview_content=generated_content,
            )

        # Full generation below
        # If idempotency upgraded a draft, gc and review_id are already set
        if 'gc' not in dir() or gc is None:
            review_id = str(uuid.uuid4())
            docx_key = f"generated/{self.repo.tenant_id}/{review_id}/contract.docx"
            pdf_key = f"generated/{self.repo.tenant_id}/{review_id}/contract.pdf"
        else:
            review_id = gc.review_id or str(uuid.uuid4())
            docx_key = f"generated/{self.repo.tenant_id}/{review_id}/contract.docx"
            pdf_key = f"generated/{self.repo.tenant_id}/{review_id}/contract.pdf"

        # Create generated contract record (skip if upgraded from draft)
        if 'gc' not in dir() or gc is None:
            gc = GeneratedContract(
                tenant_id=self.repo.tenant_id,
                template_id=body.template_id,
                template_version_id=current_ver.id,
                review_id=review_id,
                title=title,
                variable_values=body.variable_values,
                clause_selections=clause_selections_data,
                idempotency_key=body.idempotency_key,
                generated_docx_key=docx_key,
                generated_pdf_key=pdf_key,
                status=body.status or "draft",
                created_by=self.actor_id,
                created_at=now,
            )
            gc = await self.repo.create_generated_contract(gc)

        # 4. Log usage
        await self.repo.increment_usage(body.template_id)
        await self.repo.log_usage(
            body.template_id, "generated", self.actor_id,
            template_version_id=current_ver.id,
            details={
                "generated_contract_id": gc.id,
                "review_id": review_id,
                "title": title,
                "clause_count": len(clause_selections_data),
                "variable_count": len(body.variable_values),
                "template_version_id": current_ver.id,
                "template_name": t.name,
                "clause_selections": clause_selections_data,
                "idempotency_key": body.idempotency_key,
            },
        )

        # 5. Create contract review record so it appears in Contract Repository
        document_version_id: str | None = None
        upload_session_id: str | None = None
        if body.status != "draft" and not body.preview_only:
            try:
                from sqlalchemy import text as sa_text
                import json as json_mod
                from sqlalchemy import select as sa_select
                from app.domains.review.models import ContractReview

                # Check if review already exists (idempotency)
                existing_check = await self._session.execute(
                    sa_select(ContractReview).where(ContractReview.review_id == uuid.UUID(review_id))
                )
                if not existing_check.scalar_one_or_none():
                    # Create upload session
                    upload_session_id = str(uuid.uuid4())
                    await self._session.execute(
                        sa_text("""
                            INSERT INTO upload_sessions (upload_id, tenant_id, user_id, filename, content_type, file_size, ingestion_state, retry_count, metadata, created_at, updated_at)
                            VALUES (CAST(:uid AS uuid), CAST(:tid AS uuid), :user_id, :fname, :ctype, 0, 'uploaded', 0, '{}'::jsonb, :now, :now)
                        """),
                        {"uid": upload_session_id, "tid": self.repo.tenant_id, "user_id": self.actor_id,
                         "fname": f"{title}.docx", "ctype": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "now": now},
                    )

                    # Build rich metadata from variable values
                    vendor_name = body.variable_values.get("VendorName") or body.variable_values.get("Vendor") or ""
                    company_name = body.variable_values.get("CompanyName") or ""
                    effective_date = body.variable_values.get("EffectiveDate") or ""
                    contract_value = body.variable_values.get("ContractValue") or body.variable_values.get("ProjectBudget") or body.variable_values.get("EngagementFee") or ""
                    jurisdiction = body.variable_values.get("Jurisdiction") or ""

                    # Store rendered document content in generated_contracts
                    if 'gc' in dir() and gc is not None:
                        await self._session.execute(
                            sa_text("UPDATE generated_contracts SET generated_content = :content WHERE id = :id"),
                            {"content": generated_content, "id": gc.id},
                        )

                    # Calculate page count from content length (~3000 chars per page)
                    total_pages = max(1, len(generated_content) // 3000)

                    # Create review with rich metadata
                    await self._session.execute(
                        sa_text("""
                            INSERT INTO contract_reviews (
                                review_id, upload_id, tenant_id, status, priority,
                                sla_breached, finding_count, redline_count, comment_count, escalation_count,
                                is_deleted, is_favorite, version, sla_status, overdue_hours,
                                created_by, created_at, updated_at, metadata
                            ) VALUES (
                                CAST(:rid AS uuid), CAST(:uid AS uuid), CAST(:tid AS uuid),
                                :status, :priority,
                                false, 0, 0, 0, 0,
                                false, false, 1, 'on_track', 0,
                                :created_by, :now, :now, CAST(:metadata AS jsonb)
                            )
                        """),
                        {"rid": review_id, "uid": upload_session_id, "tid": self.repo.tenant_id,
                         "status": "draft", "priority": "normal", "created_by": self.actor_id,
                         "now": now, "metadata": json_mod.dumps({
                             "title": title,
                             "source": "template_generation",
                             "template_id": body.template_id,
                             "template_name": t.name,
                             "template_version_id": current_ver.id,
                             "generated_contract_id": gc.id,
                             "clause_count": len(clause_selections_data),
                             "vendor": vendor_name,
                             "name": company_name or title,
                             "effective_date": effective_date,
                             "financial_value": contract_value,
                             "currency": body.variable_values.get("Currency", "USD"),
                             "jurisdiction": jurisdiction,
                             "contract_type": t.name,
                             "owner": self.actor_id,
                             "total_pages": total_pages,
                             "generated_content_preview": generated_content[:500] if generated_content else "",
                             "analysis_status": "ingestion_queued",
                         })},
                    )

                    # Upload enterprise-grade DOCX artifact and link to upload session
                    docx_storage_key, docx_file_size = await self._upload_generated_docx(
                        review_id=review_id,
                        upload_session_id=upload_session_id,
                        content=generated_content,
                        filename=f"{title}.docx",
                        now=now,
                    )

                    # v1 document version — gold-standard generated contract
                    document_version_id = str(uuid.uuid4())
                    await self._session.execute(
                        sa_text("""
                            INSERT INTO contract_document_versions (version_id, review_id, tenant_id, version_number, label, status, storage_key, file_size_bytes, mime_type, created_by, created_at)
                            VALUES (CAST(:vid AS uuid), CAST(:rid AS uuid), CAST(:tid AS uuid), 1, 'Generated from Template', 'current', :storage_key, :file_size, :mime_type, :created_by, :now)
                        """),
                        {"vid": document_version_id, "rid": review_id, "tid": self.repo.tenant_id,
                         "storage_key": docx_storage_key,
                         "file_size": docx_file_size,
                         "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         "created_by": self.actor_id, "now": now},
                    )

                    # ── Governance audit trail for excluded required clauses ──
                    try:
                        governance_count = 0
                        avail_clauses = await self.get_available_clauses(body.template_id)
                        clause_lookup = {ac.get("clause_id"): ac for ac in avail_clauses}
                        for cs in clause_selections_data:
                            if cs.get("included", True):
                                continue
                            ref = clause_lookup.get(cs.get("clause_id"), {})
                            if not ref.get("is_required") and not cs.get("is_required"):
                                continue
                            clause_title = cs.get("clause_title") or ref.get("clause_title") or cs.get("clause_id", "Clause")
                            await self._session.execute(
                                sa_text("""
                                    INSERT INTO review_findings (
                                        finding_id, review_id, upload_id, tenant_id,
                                        clause_type, severity, title, description,
                                        confidence, risk_score, page_numbers, chunk_ids,
                                        created_at
                                    ) VALUES (
                                        CAST(:fid AS uuid), CAST(:rid AS uuid), CAST(:uid AS uuid), CAST(:tid AS uuid),
                                        :clause_type, :severity, :title, :description,
                                        :confidence, :risk_score, :page_numbers, :chunk_ids, :now
                                    )
                                """),
                                {"fid": str(uuid.uuid4()), "rid": review_id,
                                 "uid": upload_session_id, "tid": self.repo.tenant_id,
                                 "clause_type": ref.get("clause_type", "governance"),
                                 "severity": "high",
                                 "title": f"Required clause excluded: {clause_title}",
                                 "description": (
                                     f"Required clause '{clause_title}' was excluded during template generation. "
                                     f"Reason: {cs.get('reason_removed') or 'Not provided'}"
                                 ),
                                 "confidence": 1.0, "risk_score": 0.8,
                                 "page_numbers": [], "chunk_ids": [], "now": now},
                            )
                            governance_count += 1

                        if governance_count:
                            await self._session.execute(
                                sa_text("""
                                    UPDATE contract_reviews
                                    SET finding_count = :findings
                                    WHERE review_id = CAST(:rid AS uuid)
                                """),
                                {"findings": governance_count, "rid": review_id},
                            )

                        # Defer Celery dispatch until after HTTP commit (see templates router).
                        self._pending_ingestion = (
                            upload_session_id,
                            str(self.repo.tenant_id),
                            self.actor_id or "system",
                        )
                        logger.info(
                            "Template contract review=%s upload=%s pending ingestion dispatch",
                            review_id,
                            upload_session_id,
                        )

                    except Exception as exc:
                        logger.warning("Failed post-generation governance/AI setup: %s", exc)

                await self._session.flush()
            except Exception as exc:
                logger.warning("Failed to create contract review: %s", exc)

        docx_download_url = (
            f"/api/v1/reviews/{review_id}/versions/{document_version_id}/download"
            if document_version_id
            else None
        )
        return GenerateContractResponse(
            generated_contract_id=gc.id,
            review_id=review_id,
            template_version_id=current_ver.id,
            title=title,
            status=gc.status,
            variable_values=body.variable_values,
            clause_selections=clause_selections_data,
            generated_docx_url=docx_download_url,
            generated_pdf_url=docx_download_url,
            document_version_id=document_version_id,
            upload_id=upload_session_id,
            created_at=now.isoformat(),
            is_draft=gc.status == "draft",
        )

    # ── Draft Resume ──────────────────────────────────────────────

    async def save_draft(
        self, body: GenerateContractRequest, step: int = 1,
    ) -> Optional[GenerateContractResponse]:
        """Save or update a draft contract without finalizing."""
        t = await self.repo.get_template(body.template_id)
        if not t:
            return None

        now = datetime.now(timezone.utc)
        title = body.title or f"{t.name} - Draft"

        clause_selections_data = []
        if body.clause_selections:
            clause_selections_data = [cs.model_dump() for cs in body.clause_selections]

        # If idempotency_key provided, try to update existing draft
        if body.idempotency_key:
            existing = await self.repo.get_generated_contract_by_idempotency(body.idempotency_key)
            if existing:
                existing.title = title
                existing.variable_values = body.variable_values
                existing.clause_selections = clause_selections_data
                existing.updated_at = now
                await self.repo.session.flush()
                return GenerateContractResponse(
                    generated_contract_id=existing.id,
                    review_id=existing.review_id or "",
                    template_version_id=existing.template_version_id,
                    title=title,
                    status="draft",
                    variable_values=body.variable_values,
                    clause_selections=clause_selections_data,
                    created_at=existing.created_at.isoformat() if existing.created_at else None,
                    is_draft=True,
                )

        # Create new draft
        gc = GeneratedContract(
            tenant_id=self.repo.tenant_id,
            template_id=body.template_id,
            template_version_id=body.template_version_id or t.current_version_id,
            title=title,
            variable_values=body.variable_values,
            clause_selections=clause_selections_data,
            idempotency_key=body.idempotency_key,
            status="draft",
            created_by=self.actor_id,
            created_at=now,
        )
        gc = await self.repo.create_generated_contract(gc)
        return GenerateContractResponse(
            generated_contract_id=gc.id,
            title=title,
            status="draft",
            variable_values=body.variable_values,
            clause_selections=clause_selections_data,
            created_at=now.isoformat(),
            is_draft=True,
        )

    async def list_drafts(self) -> list[dict]:
        """List all drafts for the current user."""
        drafts = await self.repo.list_drafts(self.actor_id)
        return [
            {
                "id": d.id,
                "template_id": d.template_id,
                "title": d.title,
                "status": d.status,
                "variable_count": len(d.variable_values) if isinstance(d.variable_values, dict) else 0,
                "clause_count": len(d.clause_selections) if isinstance(d.clause_selections, list) else 0,
                "created_at": d.created_at.isoformat() if d.created_at else None,
                "updated_at": d.updated_at.isoformat() if d.updated_at else None,
            }
            for d in drafts
        ]

    async def get_draft(self, draft_id: str) -> Optional[dict]:
        """Get a draft by ID for resume."""
        gc = await self.repo.get_generated_contract(draft_id)
        if not gc or gc.status != "draft":
            return None
        return {
            "id": gc.id,
            "template_id": gc.template_id,
            "template_version_id": gc.template_version_id,
            "title": gc.title,
            "status": gc.status,
            "variable_values": gc.variable_values,
            "clause_selections": gc.clause_selections,
            "created_by": gc.created_by,
            "created_at": gc.created_at.isoformat() if gc.created_at else None,
            "updated_at": gc.updated_at.isoformat() if gc.updated_at else None,
        }

    # ── Clause Library ────────────────────────────────────────────

    async def list_clauses(
        self,
        clause_type: Optional[str] = None,
        template_id: Optional[str] = None,
        page: int = 1,
        page_size: int = 50,
    ) -> tuple[list[TemplateClauseResponse], int]:
        clauses, total = await self.repo.list_clauses(
            clause_type=clause_type, template_id=template_id,
            page=page, page_size=page_size,
        )
        items = [
            TemplateClauseResponse(
                id=c.id, template_id=c.template_id,
                clause_type=c.clause_type, title=c.title,
                content=c.content, category=c.category,
                risk_level=c.risk_level,
                ai_rewrite_allowed=c.ai_rewrite_allowed,
                fallback_clause_id=c.fallback_clause_id,
                is_required=c.is_required,
                is_conditional=c.is_conditional,
                condition_expression=c.condition_expression,
                display_order=c.display_order,
                status=c.status, version=c.version,
                change_summary=c.change_summary,
                approved_by=c.approved_by,
                approved_at=c.approved_at.isoformat() if c.approved_at else None,
                usage_count=c.usage_count or 0,
                contract_usage_count=c.contract_usage_count or 0,
                tags=list(c.tags) if isinstance(c.tags, list) else [],
                created_by=c.created_by, updated_by=c.updated_by,
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in clauses
        ]
        return items, total

    async def create_clause(self, body: TemplateClauseCreate) -> TemplateClauseResponse:
        clause = TemplateClause(
            tenant_id=self.repo.tenant_id,
            clause_type=body.clause_type,
            title=body.title,
            content=body.content,
            category=body.category,
            risk_level=body.risk_level,
            ai_rewrite_allowed=body.ai_rewrite_allowed,
            fallback_clause_id=body.fallback_clause_id,
            is_required=body.is_required,
            is_conditional=body.is_conditional,
            condition_expression=body.condition_expression,
            display_order=body.display_order,
            change_summary=body.change_summary,
            tags=body.tags or [],
            created_by=self.actor_id,
        )
        clause = await self.repo.create_clause(clause)
        return await self.get_clause(clause.id)

    async def get_clause(self, clause_id: str) -> Optional[TemplateClauseResponse]:
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None
        return TemplateClauseResponse(
            id=clause.id, template_id=clause.template_id,
            clause_type=clause.clause_type, title=clause.title,
            content=clause.content, category=clause.category,
            risk_level=clause.risk_level,
            ai_rewrite_allowed=clause.ai_rewrite_allowed,
            fallback_clause_id=clause.fallback_clause_id,
            is_required=clause.is_required,
            is_conditional=clause.is_conditional,
            condition_expression=clause.condition_expression,
            display_order=clause.display_order,
            status=clause.status, version=clause.version,
            change_summary=clause.change_summary,
            approved_by=clause.approved_by,
            approved_at=clause.approved_at.isoformat() if clause.approved_at else None,
            usage_count=clause.usage_count or 0,
            contract_usage_count=clause.contract_usage_count or 0,
            tags=list(clause.tags) if isinstance(clause.tags, list) else [],
            created_by=clause.created_by, updated_by=clause.updated_by,
            created_at=clause.created_at.isoformat() if clause.created_at else None,
            updated_at=clause.updated_at.isoformat() if clause.updated_at else None,
        )

    async def update_clause(self, clause_id: str, body: TemplateClauseUpdate) -> Optional[TemplateClauseResponse]:
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None

        # Track if content changed — triggers version bump
        content_changed = False

        if body.title is not None:
            clause.title = body.title
        if body.content is not None:
            clause.content = body.content
            content_changed = True
        if body.category is not None:
            clause.category = body.category
        if body.risk_level is not None:
            clause.risk_level = body.risk_level
        if body.ai_rewrite_allowed is not None:
            clause.ai_rewrite_allowed = body.ai_rewrite_allowed
        if body.fallback_clause_id is not None:
            clause.fallback_clause_id = body.fallback_clause_id
        if body.is_required is not None:
            clause.is_required = body.is_required
        if body.is_conditional is not None:
            clause.is_conditional = body.is_conditional
        if body.condition_expression is not None:
            clause.condition_expression = body.condition_expression
        if body.display_order is not None:
            clause.display_order = body.display_order
        if body.status is not None:
            clause.status = body.status
        if body.change_summary is not None:
            clause.change_summary = body.change_summary
        if body.tags is not None:
            clause.tags = body.tags

        # Bump version on content change
        if content_changed:
            clause.version += 1
            # Reset to draft when content changes
            if clause.status in ("approved", "published"):
                clause.status = "draft"
                clause.approved_by = None
                clause.approved_at = None

        clause.updated_by = self.actor_id
        await self.repo.update_clause(clause)
        return await self.get_clause(clause.id)

    async def delete_clause(self, clause_id: str) -> bool:
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return False
        await self.repo.delete_clause(clause)
        return True

    # ── Clause Approval Workflow ──────────────────────────────────

    async def submit_clause_for_review(self, clause_id: str) -> Optional[TemplateClauseResponse]:
        """Submit a clause from draft → legal_review."""
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None
        if clause.status != "draft":
            raise ValueError(f"Cannot submit clause with status '{clause.status}'. Only draft clauses can be submitted.")
        clause.status = "legal_review"
        clause.updated_by = self.actor_id
        await self.repo.update_clause(clause)
        return await self.get_clause(clause.id)

    async def approve_clause(self, clause_id: str) -> Optional[TemplateClauseResponse]:
        """Approve a clause — legal_review → approved."""
        from datetime import datetime, timezone
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None
        if clause.status != "legal_review":
            raise ValueError(f"Cannot approve clause with status '{clause.status}'. Only clauses in legal_review can be approved.")
        clause.status = "approved"
        clause.approved_by = self.actor_id
        clause.approved_at = datetime.now(timezone.utc)
        clause.updated_by = self.actor_id
        await self.repo.update_clause(clause)
        return await self.get_clause(clause.id)

    async def publish_clause(self, clause_id: str) -> Optional[TemplateClauseResponse]:
        """Publish a clause — approved → published."""
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None
        if clause.status != "approved":
            raise ValueError(f"Cannot publish clause with status '{clause.status}'. Only approved clauses can be published.")
        clause.status = "published"
        clause.updated_by = self.actor_id
        await self.repo.update_clause(clause)
        return await self.get_clause(clause.id)

    async def reject_clause(self, clause_id: str, reason: str = "") -> Optional[TemplateClauseResponse]:
        """Reject a clause — send back to draft with feedback."""
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None
        if clause.status not in ("legal_review", "approved"):
            raise ValueError(f"Cannot reject clause with status '{clause.status}'.")
        clause.status = "draft"
        clause.approved_by = None
        clause.approved_at = None
        clause.change_summary = f"Rejected: {reason}" if reason else "Rejected — returned to draft"
        clause.updated_by = self.actor_id
        await self.repo.update_clause(clause)
        return await self.get_clause(clause.id)

    # ── Clause Dependency Report ──────────────────────────────────

    async def check_clause_dependencies(self, clause_id: str) -> Optional[ClauseDependencyInfo]:
        """Check how many templates and contracts reference a clause."""
        clause = await self.repo.get_clause(clause_id)
        if not clause:
            return None

        deps = await self.repo.get_clause_dependencies(clause_id)
        blocking_reasons: list[str] = []

        if deps["template_count"] > 0:
            blocking_reasons.append(f"Used by {deps['template_count']} template(s)")
        if deps["contract_count"] > 0:
            blocking_reasons.append(f"Used by {deps['contract_count']} contract(s)")

        return ClauseDependencyInfo(
            clause_id=clause_id,
            clause_title=clause.title,
            template_count=deps["template_count"],
            contract_count=deps["contract_count"],
            templates=deps["templates"],
            can_delete=deps["template_count"] == 0 and deps["contract_count"] == 0,
            blocking_reasons=blocking_reasons,
        )

    # ── Clause Usage Analytics ────────────────────────────────────

    async def get_clause_analytics(self) -> ClauseAnalytics:
        """Get clause library analytics — most used, highest risk, deprecated, etc."""
        analytics = await self.repo.get_clause_analytics()

        most_used = [
            TemplateClauseResponse(
                id=c.id, template_id=c.template_id,
                clause_type=c.clause_type, title=c.title,
                content=c.content, category=c.category,
                risk_level=c.risk_level,
                ai_rewrite_allowed=c.ai_rewrite_allowed,
                fallback_clause_id=c.fallback_clause_id,
                is_required=c.is_required,
                is_conditional=c.is_conditional,
                condition_expression=c.condition_expression,
                display_order=c.display_order,
                status=c.status, version=c.version,
                usage_count=c.usage_count,
                contract_usage_count=c.contract_usage_count,
                created_by=c.created_by, updated_by=c.updated_by,
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in analytics.get("most_used_clauses", [])
        ]

        highest_risk = [
            TemplateClauseResponse(
                id=c.id, template_id=c.template_id,
                clause_type=c.clause_type, title=c.title,
                content=c.content, category=c.category,
                risk_level=c.risk_level,
                ai_rewrite_allowed=c.ai_rewrite_allowed,
                fallback_clause_id=c.fallback_clause_id,
                is_required=c.is_required,
                is_conditional=c.is_conditional,
                condition_expression=c.condition_expression,
                display_order=c.display_order,
                status=c.status, version=c.version,
                usage_count=c.usage_count,
                contract_usage_count=c.contract_usage_count,
                created_by=c.created_by, updated_by=c.updated_by,
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in analytics.get("highest_risk_clauses", [])
        ]

        deprecated_still_used = [
            TemplateClauseResponse(
                id=c.id, template_id=c.template_id,
                clause_type=c.clause_type, title=c.title,
                content=c.content, category=c.category,
                risk_level=c.risk_level,
                ai_rewrite_allowed=c.ai_rewrite_allowed,
                fallback_clause_id=c.fallback_clause_id,
                is_required=c.is_required,
                is_conditional=c.is_conditional,
                condition_expression=c.condition_expression,
                display_order=c.display_order,
                status=c.status, version=c.version,
                usage_count=c.usage_count,
                contract_usage_count=c.contract_usage_count,
                created_by=c.created_by, updated_by=c.updated_by,
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in analytics.get("deprecated_still_used", [])
        ]

        return ClauseAnalytics(
            total_clauses=analytics.get("total_clauses", 0),
            published_clauses=analytics.get("published_clauses", 0),
            draft_clauses=analytics.get("draft_clauses", 0),
            most_used_clauses=most_used,
            highest_risk_clauses=highest_risk,
            deprecated_still_used=deprecated_still_used,
            avg_ai_rewrite_rate=analytics.get("avg_ai_rewrite_rate", 0),
            clauses_by_type=analytics.get("clauses_by_type", []),
            clauses_by_risk=analytics.get("clauses_by_risk", []),
        )

    # ── Clause Search (for Global Search integration) ─────────────

    async def search_clauses_global(
        self,
        query: str,
        clause_type: Optional[str] = None,
        risk_level: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[list[TemplateClauseResponse], int]:
        """Search clauses by name, content, tags, risk, and category.

        Used by the Global Search feature to include clause library results.
        """
        clauses, total = await self.repo.search_clauses(
            query=query, clause_type=clause_type,
            risk_level=risk_level, page=page, page_size=page_size,
        )
        items = [
            TemplateClauseResponse(
                id=c.id, template_id=c.template_id,
                clause_type=c.clause_type, title=c.title,
                content=c.content, category=c.category,
                risk_level=c.risk_level,
                ai_rewrite_allowed=c.ai_rewrite_allowed,
                fallback_clause_id=c.fallback_clause_id,
                is_required=c.is_required,
                is_conditional=c.is_conditional,
                condition_expression=c.condition_expression,
                display_order=c.display_order,
                status=c.status, version=c.version,
                usage_count=c.usage_count or 0,
                contract_usage_count=c.contract_usage_count or 0,
                tags=list(c.tags) if hasattr(c, 'tags') and isinstance(c.tags, list) else [],
                created_by=c.created_by, updated_by=c.updated_by,
                created_at=c.created_at.isoformat() if c.created_at else None,
                updated_at=c.updated_at.isoformat() if c.updated_at else None,
            )
            for c in clauses
        ]
        return items, total

    # ── Clause Refs ───────────────────────────────────────────────

    async def set_template_clause_refs(
        self, template_version_id: str, refs: list[dict],
    ) -> list[dict]:
        """Set clause references for a template version, pinning clause versions."""
        return await self.repo.set_clause_refs(template_version_id, refs)

    async def get_template_clause_refs(
        self, template_version_id: str,
    ) -> list[dict]:
        """Get clause references for a template version with clause details."""
        return await self.repo.get_clause_refs(template_version_id)

    # ── Template Packages ────────────────────────────────────────

    async def list_packages(
        self,
        industry: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[list[TemplatePackageResponse], int]:
        packages, total = await self.repo.list_packages(
            industry=industry, page=page, page_size=page_size,
        )
        items = []
        for p in packages:
            items.append(await self._package_to_response(p))
        return items, total

    async def get_package(self, package_id: str) -> Optional[TemplatePackageResponse]:
        pkg = await self.repo.get_package(package_id)
        if not pkg:
            return None
        return await self._package_to_response(pkg)

    async def create_package(self, body: TemplatePackageCreate) -> TemplatePackageResponse:
        now = datetime.now(timezone.utc)
        pkg = TemplatePackage(
            tenant_id=self.repo.tenant_id,
            name=body.name,
            industry=body.industry,
            description=body.description,
            tags=body.tags or [],
            icon=body.icon,
            created_by=self.actor_id,
            created_at=now,
            updated_at=now,
        )
        pkg = await self.repo.create_package(pkg)

        # Add items if provided
        if body.items:
            for i, item in enumerate(body.items):
                pi = TemplatePackageItem(
                    tenant_id=self.repo.tenant_id,
                    package_id=pkg.id,
                    template_id=item.template_id,
                    display_order=item.display_order if item.display_order is not None else i,
                    is_required=item.is_required,
                )
                self.repo.session.add(pi)
            await self.repo.session.flush()

        return await self.get_package(pkg.id)

    async def update_package(
        self, package_id: str, body: TemplatePackageUpdate,
    ) -> Optional[TemplatePackageResponse]:
        pkg = await self.repo.get_package(package_id)
        if not pkg:
            return None
        if body.name is not None:
            pkg.name = body.name
        if body.industry is not None:
            pkg.industry = body.industry
        if body.description is not None:
            pkg.description = body.description
        if body.tags is not None:
            pkg.tags = body.tags
        if body.icon is not None:
            pkg.icon = body.icon
        if body.is_published is not None:
            pkg.is_published = body.is_published
        pkg.updated_by = self.actor_id
        pkg.version += 1
        await self.repo.update_package(pkg)
        return await self.get_package(package_id)

    async def delete_package(self, package_id: str) -> bool:
        return await self.repo.delete_package(package_id)

    async def publish_package(self, package_id: str) -> Optional[TemplatePackageResponse]:
        return await self.update_package(
            package_id, TemplatePackageUpdate(is_published=True),
        )

    async def set_package_items(
        self, package_id: str, items: list[PackageItemCreate],
    ) -> Optional[list[PackageItemResponse]]:
        pkg = await self.repo.get_package(package_id)
        if not pkg:
            return None
        result = await self.repo.set_package_items(package_id, items)
        return result

    async def _package_to_response(self, pkg: TemplatePackage) -> TemplatePackageResponse:
        items = []
        for item in (pkg.items or []):
            template_name = None
            template_status = None
            try:
                tpl = await self.repo.get_template(item.template_id)
                if tpl:
                    template_name = tpl.name
                    template_status = tpl.status
            except Exception:
                pass
            items.append(PackageItemResponse(
                id=item.id,
                package_id=item.package_id,
                template_id=item.template_id,
                display_order=item.display_order,
                is_required=item.is_required,
                template_name=template_name,
                template_status=template_status,
            ))
        return TemplatePackageResponse(
            id=pkg.id,
            name=pkg.name,
            industry=pkg.industry,
            description=pkg.description,
            version=pkg.version,
            is_published=pkg.is_published,
            tags=list(pkg.tags) if isinstance(pkg.tags, list) else [],
            icon=pkg.icon,
            items=items,
            template_count=len(items),
            created_by=pkg.created_by,
            updated_by=pkg.updated_by,
            created_at=pkg.created_at.isoformat() if pkg.created_at else None,
            updated_at=pkg.updated_at.isoformat() if pkg.updated_at else None,
        )

    async def _assemble_generated_content(
        self,
        template_content: str,
        variable_values: dict[str, Any],
        clause_selections: list[dict],
        template_version_id: str,
    ) -> str:
        """Build the full contract body: template + substituted variables + included clauses."""
        body = self._substitute_variables(template_content, variable_values)
        if not clause_selections:
            return body

        available = await self.repo.get_clause_refs(template_version_id)
        clause_by_id = {c.get("clause_id"): c for c in available}
        sections: list[str] = []

        for cs in sorted(clause_selections, key=lambda x: x.get("sort_order", 0)):
            if not cs.get("included", True):
                continue
            ref = clause_by_id.get(cs.get("clause_id"), {})
            title = cs.get("clause_title") or ref.get("clause_title") or cs.get("clause_id", "Clause")
            content = cs.get("content") or ref.get("clause_content") or ""
            if content:
                content = self._substitute_variables(content, variable_values)
                sections.append(f"\n\n## {title}\n\n{content.strip()}")

        if sections:
            body = body.rstrip() + "".join(sections)
        return body

    async def _upload_generated_docx(
        self,
        review_id: str,
        upload_session_id: str,
        content: str,
        filename: str,
        now: datetime,
    ) -> tuple[str, int]:
        """Create and upload a formatted DOCX for a template-generated contract."""
        import hashlib
        import io
        from sqlalchemy import text as sa_text
        from docx import Document as DocxDocument
        from app.integrations.storage.s3 import storage_service
        from app.config import settings

        doc = DocxDocument()
        doc.add_heading(filename.replace(".docx", ""), level=0)
        for line in content.split("\n"):
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph("")
                continue
            if stripped.startswith("## "):
                doc.add_heading(stripped[3:].strip(), level=2)
            elif stripped.isupper() and len(stripped) > 3:
                doc.add_heading(stripped, level=1)
            else:
                doc.add_paragraph(line)

        output = io.BytesIO()
        doc.save(output)
        docx_bytes = output.getvalue()

        storage_key = f"generated/{self.repo.tenant_id}/{review_id}/contract.docx"
        checksum = hashlib.sha256(docx_bytes).hexdigest()
        bucket = settings.s3_bucket or "contractrisk-documents"
        await storage_service.upload_fileobj(
            bucket=bucket,
            key=storage_key,
            file_body=docx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            metadata={
                "review_id": review_id,
                "source": "template_generation",
                "checksum_sha256": checksum,
            },
        )

        await self._session.execute(
            sa_text("""
                UPDATE upload_sessions
                SET storage_key = :storage_key,
                    file_size = :file_size,
                    content_type = :ctype,
                    server_checksum_sha256 = :checksum,
                    ingestion_state = 'uploaded',
                    updated_at = :now
                WHERE upload_id = CAST(:uid AS uuid)
            """),
            {
                "storage_key": storage_key,
                "file_size": len(docx_bytes),
                "ctype": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "checksum": checksum,
                "uid": upload_session_id,
                "now": now,
            },
        )
        return storage_key, len(docx_bytes)

    def _substitute_variables(
        self, content: str, values: dict[str, Any],
    ) -> str:
        """Replace {{VariableName}} placeholders with actual values."""
        def replacer(match):
            key = match.group(1).strip()
            val = values.get(key, match.group(0))
            if val is None:
                return ""
            return str(val)

        return re.sub(r"\{\{(\w+)\}\}", replacer, content)

    @staticmethod
    def extract_variables(content: str) -> list[dict]:
        """Extract all {{VariableName}} placeholders from template content."""
        seen = set()
        variables = []
        for match in re.finditer(r"\{\{(\w+)\}\}", content):
            key = match.group(1)
            if key not in seen:
                seen.add(key)
                # Generate a human-readable label from camelCase/PascalCase
                label = re.sub(r"([a-z])([A-Z])", r"\1 \2", key)
                label = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", label)
                variables.append({
                    "key": key,
                    "label": label,
                    "field_type": self._infer_field_type(key),
                    "is_required": True,
                    "display_order": len(seen),
                })
        return variables

    @staticmethod
    def _infer_field_type(key: str) -> str:
        """Infer the most likely field type from a variable name."""
        key_lower = key.lower()
        if any(word in key_lower for word in ["date", "day", "effective", "expiration", "renewal"]):
            return "date"
        if any(word in key_lower for word in ["email"]):
            return "email"
        if any(word in key_lower for word in ["phone", "telephone", "fax"]):
            return "phone"
        if any(word in key_lower for word in ["url", "website", "web"]):
            return "url"
        if any(word in key_lower for word in ["value", "amount", "price", "fee", "cost", "budget"]):
            return "currency"
        if any(word in key_lower for word in ["count", "number", "quantity", "percent", "rate"]):
            return "number"
        if any(word in key_lower for word in ["boolean", "is_", "has_", "enable", "flag"]):
            return "boolean"
        if any(word in key_lower for word in ["address", "location"]):
            return "address"
        return "text"

    # ── Dashboard Metrics ─────────────────────────────────────────

    async def get_metrics(self) -> TemplateMetrics:
        metrics = await self.repo.get_metrics()
        most_used_items = []
        for t in metrics.get("most_used_templates", []):
            most_used_items.append(TemplateListItem(
                id=t.id, name=t.name, description=t.description,
                status=t.status, usage_count=t.usage_count,
                created_by=t.created_by, is_favorite=False,
            ))
        return TemplateMetrics(
            total_templates=metrics.get("total_templates", 0),
            approved_templates=metrics.get("approved_templates", 0),
            draft_templates=metrics.get("draft_templates", 0),
            generated_this_month=metrics.get("generated_this_month", 0),
            most_used_templates=most_used_items,
            top_categories=metrics.get("top_categories", []),
        )

    # ── Template Validation ───────────────────────────────────────

    async def validate_template(self, template_id: str) -> TemplateValidationResult:
        """Validate a template before publishing — checks placeholders, variables, clauses."""
        t = await self.repo.get_template(template_id)
        if not t:
            raise ValueError("Template not found")

        current_ver = None
        for v in (t.versions or []):
            if v.id == t.current_version_id:
                current_ver = v
                break
        if not current_ver and t.versions:
            current_ver = t.versions[0]
        if not current_ver:
            return TemplateValidationResult(
                is_valid=False,
                issues=[ValidationIssue(severity="error", message="Template has no versions")],
            )

        issues: list[ValidationIssue] = []
        content = current_ver.placeholder_content or ""
        variables = current_ver.variables or []

        # Extract all placeholders from content
        placeholders = set(re.findall(r"\{\{(\w+)\}\}", content))

        # Check for duplicate placeholders in content
        all_placeholders = re.findall(r"\{\{(\w+)\}\}", content)
        seen = set()
        duplicates = set()
        for p in all_placeholders:
            if p in seen:
                duplicates.add(p)
            seen.add(p)

        # Check mapped variables
        mapped_keys = {v.get("key") for v in variables if isinstance(v, dict)}
        unused_vars = mapped_keys - placeholders
        missing_placeholders = placeholders - mapped_keys

        if duplicates:
            issues.append(ValidationIssue(
                severity="warning",
                message=f"Duplicate placeholders found: {', '.join(duplicates)}",
            ))

        if unused_vars:
            issues.append(ValidationIssue(
                severity="warning",
                message=f"Unused variables: {', '.join(unused_vars)}",
            ))

        if missing_placeholders:
            issues.append(ValidationIssue(
                severity="error",
                message=f"Missing variable mappings: {', '.join(missing_placeholders)}",
            ))

        # Check for invalid placeholder syntax
        invalid_syntax = re.findall(r"\{\{([^}]*\s+[^}]*)\}\}", content)
        if invalid_syntax:
            issues.append(ValidationIssue(
                severity="error",
                message=f"Invalid placeholder syntax: {', '.join(invalid_syntax[:3])}",
            ))

        # ── Enterprise validation checks ──────────────────────────

        # 1. Missing required variables (required fields with no value/default)
        for v in variables:
            if isinstance(v, dict) and v.get("is_required"):
                key = v.get("key", "")
                default_val = v.get("default_value")
                if not default_val and key in placeholders:
                    issues.append(ValidationIssue(
                        severity="warning",
                        field=key,
                        message=f"Required variable '{v.get('label', key)}' has no default value",
                    ))

        # 2. Missing required clauses
        clause_refs = current_ver.clause_refs or []
        for ref in clause_refs:
            if isinstance(ref, dict) and ref.get("is_required"):
                clause_id = ref.get("clause_id", "")
                issues.append(ValidationIssue(
                    severity="info",
                    field=clause_id,
                    message=f"Required clause '{ref.get('clause_title', clause_id)}' is included by default",
                ))

        # 3. Duplicate clause references
        clause_ids = [ref.get("clause_id") for ref in clause_refs if isinstance(ref, dict)]
        seen_clauses = set()
        dup_clauses = set()
        for cid in clause_ids:
            if cid in seen_clauses:
                dup_clauses.add(cid)
            seen_clauses.add(cid)
        if dup_clauses:
            issues.append(ValidationIssue(
                severity="error",
                message=f"Duplicate clause references: {', '.join(dup_clauses)}",
            ))

        # 4. Expired clauses (clause version is behind latest)
        if clause_refs:
            try:
                from app.domains.templates.models import TemplateClause
                from sqlalchemy import select
                for ref in clause_refs:
                    if isinstance(ref, dict) and ref.get("clause_id"):
                        clause_result = await self.repo.session.execute(
                            select(TemplateClause).where(
                                TemplateClause.id == ref["clause_id"],
                                TemplateClause.tenant_id == self.repo.tenant_id,
                            )
                        )
                        clause = clause_result.scalar_one_or_none()
                        if clause and clause.status == "deprecated":
                            issues.append(ValidationIssue(
                                severity="error",
                                message=f"Deprecated clause '{clause.title}' is referenced",
                            ))
                        elif clause and ref.get("clause_version", 1) < clause.version:
                            issues.append(ValidationIssue(
                                severity="warning",
                                message=f"Clause '{clause.title}' has a newer version (v{clause.version})",
                            ))
            except Exception:
                pass

        # 5. Clause dependency missing — check fallback_clause_id references
        for ref in clause_refs:
            if isinstance(ref, dict) and ref.get("fallback_clause_id"):
                fallback_id = ref["fallback_clause_id"]
                fallback_exists = any(
                    r.get("clause_id") == fallback_id
                    for r in clause_refs if isinstance(r, dict)
                )
                if not fallback_exists:
                    issues.append(ValidationIssue(
                        severity="warning",
                        message=f"Clause '{ref.get('clause_title', ref['clause_id'])}' references missing fallback clause",
                    ))

        # 6. Deprecated clauses — checked above alongside expired

        # 7. Placeholder left empty — check for placeholders with no variable mapping
        if missing_placeholders:
            for mp in missing_placeholders:
                issues.append(ValidationIssue(
                    severity="error",
                    field=mp,
                    message=f"Placeholder '{{{{{mp}}}}}' has no variable mapping",
                ))

        is_valid = all(i.severity == "warning" for i in issues)

        return TemplateValidationResult(
            is_valid=is_valid,
            issues=issues,
            placeholder_count=len(placeholders),
            mapped_count=len(mapped_keys),
            duplicate_count=len(duplicates),
            unused_variables=list(unused_vars),
            missing_placeholders=list(missing_placeholders),
        )

    # ── Dependency Check ──────────────────────────────────────────

    async def check_dependencies(self, template_id: str) -> TemplateDependencyInfo:
        """Check if a template can be archived or deleted based on usage."""
        t = await self.repo.get_template(template_id)
        if not t:
            raise ValueError("Template not found")

        blocking_reasons: list[str] = []
        generated_count = 0
        last_used = None

        try:
            from sqlalchemy import text as sa_text
            result = await self.repo.session.execute(
                sa_text("""
                    SELECT COUNT(*)::int AS cnt, MAX(created_at) AS last_used
                    FROM generated_contracts
                    WHERE template_id = :tid AND tenant_id = :tenant_id
                """),
                {"tid": template_id, "tenant_id": self.repo.tenant_id},
            )
            row = result.fetchone()
            if row:
                generated_count = row.cnt or 0
                last_used = row.last_used.isoformat() if row.last_used else None
        except Exception:
            pass

        if generated_count > 0:
            blocking_reasons.append(
                f"Template has been used to generate {generated_count} contract(s)"
            )

        can_archive = True
        can_delete = generated_count == 0

        return TemplateDependencyInfo(
            template_id=template_id,
            template_name=t.name,
            generated_contract_count=generated_count,
            last_used=last_used,
            can_archive=can_archive,
            can_delete=can_delete,
            blocking_reasons=blocking_reasons,
        )

    # ── Generate Preview ──────────────────────────────────────────

    async def generate_preview(
        self, body: GenerateContractRequest,
    ) -> Optional[GenerateContractResponse]:
        """Generate a preview of a contract without creating any records."""
        t = await self.repo.get_template(body.template_id)
        if not t:
            return None

        current_ver = None
        for v in (t.versions or []):
            if v.id == t.current_version_id:
                current_ver = v
                break
        if not current_ver and t.versions:
            current_ver = t.versions[0]
        if not current_ver:
            raise ValueError("Template has no versions")

        content = current_ver.placeholder_content or ""
        clause_selections_data = [cs.model_dump() for cs in body.clause_selections] if body.clause_selections else []
        rendered = await self._assemble_generated_content(
            content,
            body.variable_values,
            clause_selections_data,
            current_ver.id,
        )
        title = body.title or f"{t.name} - Preview"

        return GenerateContractResponse(
            title=title,
            status="preview",
            variable_values=body.variable_values,
            preview_content=rendered,
        )

    # ── Available Clauses for Generation ──────────────────────────

    async def get_available_clauses(self, template_id: str) -> list[dict]:
        """Get all clauses available for a template when generating a contract.

        Returns clause_refs (pinned to the current version) enriched with
        full clause details from the clause library.
        """
        t = await self.repo.get_template(template_id)
        if not t:
            return []

        current_ver = None
        for v in (t.versions or []):
            if v.id == t.current_version_id:
                current_ver = v
                break
        if not current_ver and t.versions:
            current_ver = t.versions[0]
        if not current_ver:
            return []

        # Get clause refs for this version
        refs = await self.repo.get_clause_refs(current_ver.id)
        return refs

    # ── Usage History ─────────────────────────────────────────────

    async def get_usage_history(
        self, template_id: str, limit: int = 50,
    ) -> list[dict]:
        entries = await self.repo.get_usage_history(template_id, limit)
        return [
            {
                "id": e.id,
                "template_id": e.template_id,
                "action": e.action,
                "actor_id": e.actor_id,
                "details": e.details,
                "created_at": e.created_at.isoformat() if e.created_at else None,
            }
            for e in entries
        ]

    # ── Variable Extraction ───────────────────────────────────────

    async def extract_variables_from_template(
        self, template_id: str,
    ) -> Optional[list[dict]]:
        """Extract all {{variables}} from the current version of a template."""
        t = await self.repo.get_template(template_id)
        if not t:
            return None
        current_ver = None
        for v in (t.versions or []):
            if v.id == t.current_version_id:
                current_ver = v
                break
        if not current_ver and t.versions:
            current_ver = t.versions[0]
        if not current_ver or not current_ver.placeholder_content:
            return []
        return self.extract_variables(current_ver.placeholder_content)
