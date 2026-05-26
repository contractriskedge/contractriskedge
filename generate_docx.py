#!/usr/bin/env python3
"""Convert enterprise-information-architecture.md to Word .docx"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()

# Style setup
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title
title = doc.add_heading('ContractRiskEdge \u2014 Enterprise Information Architecture', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('Fortune 500 AI-First Contract Intelligence Platform')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in sub.runs:
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(100, 116, 139)

doc.add_paragraph()

with open('docs/enterprise-information-architecture.md', 'r') as f:
    content = f.read()

lines = content.split('\n')
i = 0
while i < len(lines):
    line = lines[i]

    if line.startswith('# ContractRiskEdge'):
        i += 1
        continue
    if line.strip() == '---':
        i += 1
        continue
    if line.strip().startswith('## Fortune 500'):
        i += 1
        continue

    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('#### '):
        doc.add_heading(line[5:].strip(), level=4)
    elif line.strip().startswith('|') and line.strip().endswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip())
            i += 1
        i -= 1

        if len(table_lines) >= 2:
            headers = [h.strip() for h in table_lines[0].split('|')[1:-1]]
            rows_data = []
            for tl in table_lines[2:]:
                cells = [c.strip() for c in tl.split('|')[1:-1]]
                if cells:
                    rows_data.append(cells)
            if headers and rows_data:
                table = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
                table.style = 'Light Grid Accent 1'
                for j, h in enumerate(headers):
                    cell = table.rows[0].cells[j]
                    cell.text = h
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.bold = True
                            r.font.size = Pt(9)
                for r_idx, row in enumerate(rows_data):
                    for c_idx, val in enumerate(row):
                        if c_idx < len(headers):
                            cell = table.rows[r_idx + 1].cells[c_idx]
                            cell.text = val
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.size = Pt(9)
                doc.add_paragraph()
    elif line.strip().startswith('- ') or line.strip().startswith('+ '):
        indent = len(line) - len(line.lstrip())
        text = line.strip()[2:].strip()
        if text:
            if indent > 2:
                doc.add_paragraph(text, style='List Bullet 2')
            else:
                doc.add_paragraph(text, style='List Bullet')
    elif line.strip().startswith('```'):
        while i < len(lines) and not lines[i].strip().startswith('```'):
            i += 1
    elif line.strip() and not line.strip().startswith('```'):
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line.strip())
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)

    i += 1

outpath = 'docs/enterprise-information-architecture.docx'
doc.save(outpath)
print(f'Word document saved to {outpath}')
